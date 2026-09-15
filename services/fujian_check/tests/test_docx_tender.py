"""Word 版招标文件：由磋商 PDF 合成 docx（文本 + 表格），验证 docx 读取器的合成分页能让页式解析器正常工作。"""
from __future__ import annotations

import asyncio
import re

import pytest

from services.fujian_check.tests.test_gov_cs_golden import BID_CS, TENDER_CS


@pytest.fixture(scope="session")
def cs_docx(tmp_path_factory):
    if not TENDER_CS.exists():
        pytest.skip("磋商黄金样本不存在")
    import fitz
    import pdfplumber
    from docx import Document

    out = tmp_path_factory.mktemp("docx") / "cs_tender.docx"
    d = Document()
    with fitz.open(str(TENDER_CS)) as pdf, pdfplumber.open(str(TENDER_CS)) as pl:
        for i, page in enumerate(pdf):
            for ln in page.get_text().split("\n"):
                ln = ln.strip()
                if not ln or re.match(r"^-第\d+页-$", ln):
                    continue
                if re.match(r"^第[一二三四五六七八九十\d]+章", ln):
                    d.add_heading(ln, level=1)
                else:
                    d.add_paragraph(ln)
            for tb in pl.pages[i].extract_tables() or []:
                rows = [[(c or "").strip() for c in r] for r in tb if r]
                if not rows:
                    continue
                t = d.add_table(rows=0, cols=max(len(r) for r in rows))
                for r in rows:
                    cells = t.add_row().cells
                    for j, c in enumerate(r):
                        cells[j].text = c
    d.save(str(out))
    return out


def test_docx_tender_structure(cs_docx):
    from services.fujian_check.docx_reader import parse_docx
    from services.fujian_check.profiles import detect_profile
    from services.fujian_check.tender.extract import extract_requirements

    doc = parse_docx(str(cs_docx))
    assert doc.kind == "docx" and any(p.is_toc for p in doc.pages)          # 目录行合并成目录页
    prof = detect_profile(doc)
    assert prof.key == "fujian_gov_cs"
    req = extract_requirements(doc, prof)
    chapters = {s.chapter: s for s in req.sections if s.section is None}
    assert set(chapters) >= {1, 2, 3, 4, 5} and chapters[3].page_start > chapters[2].page_end
    assert len(req.qianfubiao) == 13
    assert len([f for f in req.forms_in("资格文件") if f.no.isdigit()]) == 11
    ids = {c.id for c in req.rejection}
    assert {f"14.2.1-{i}" for i in range(1, 8)} <= ids and {f"★技术{i}" for i in range(1, 5)} <= ids
    assert len(req.datasheet) == 12 and req.staffing.roles["施工员"].count == 2
    assert req.hard.control_price.value == 1_796_580 and req.hard.validity_days == 90 and req.hard.duration_days == 90


def test_docx_tender_run(cs_docx):
    if not BID_CS.exists():
        pytest.skip("响应文件样本不存在")
    from services.fujian_check.engine import run_check

    r = asyncio.run(run_check(str(cs_docx), str(BID_CS), llm=None, use_llm=False, use_ocr=False))
    assert r.summary["total"] >= 45
    assert any("代理服务费" in f.requirement and f.verdict == "fail" for f in r.findings)
    assert any(f.rule_id == "GC-V-01" and "最高限价" in f.requirement and f.verdict == "pass" for f in r.findings)
